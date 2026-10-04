"""
Genera el informe de la evidencia GA4-210602033-AA1-EV03 (bodega de datos Spring Valley)
en Word (.docx, formato APA 7) y en PDF (HTML impreso con Chromium), desde un mismo contenido.
Las figuras son capturas reales de capturas/ (CloudBeaver, SQL Developer, logs de SQLcl).
"""
import base64, html, os, sys
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, Cm
from PIL import Image

RAIZ = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
CAP = os.path.join(RAIZ, "capturas")
CHROMIUM = "/opt/pw-browsers/chromium-1194/chrome-linux/chrome"
NOMBRE = "GA4-210602033-AA1-EV03_Bodega_Datos_Spring_Valley"

B = []  # bloques: (tipo, datos)
def h1(t): B.append(("h1", t))
def h2(t): B.append(("h2", t))
def p(t): B.append(("p", t))
def li(*items): B.append(("ul", items))
def code(t): B.append(("code", t.strip("\n")))
def tabla(titulo, enc, filas, anchos=None): B.append(("tabla", (titulo, enc, filas, anchos)))
def fig(archivo, titulo, texto): B.append(("fig", (archivo, titulo, texto)))
def salto(): B.append(("salto", None))

# ============================ CONTENIDO ============================
h1("Introducción")
p("Las empresas que crecen apoyadas en su infraestructura de tecnología acumulan grandes volúmenes de datos en "
  "sistemas diseñados para operar, no para analizar. En el caso de estudio de esta evidencia, la empresa Spring Valley "
  "sufrió un colapso de su sistema transaccional durante la facturación, cuando se intentaron ejecutar sobre él "
  "consultas históricas complejas para la toma de decisiones de la dirección. La solución propuesta es construir una "
  "bodega de datos (data warehouse) para el área de ventas, separada del sistema operativo.")
p("Este documento es una guía paso a paso para el diseño e implementación de esa bodega. Se aplica la metodología "
  "Kimball explicada en las sesiones: selección del proceso de negocio, declaración del grano, identificación de "
  "dimensiones y hechos, y construcción de un modelo estrella con una tabla de hechos FACT_VENTAS y cinco dimensiones. "
  "Además, se ejecutó realmente el proceso ETL sobre un archivo de ventas crudas con errores intencionales, se "
  "validó la carga y se respondieron con SQL las preguntas del negocio.")
p("Toda la implementación se realizó en un entorno cloud (máquina virtual Linux), porque el equipo personal del "
  "aprendiz no tiene los recursos necesarios para Oracle. Se usó Oracle Database Free en un contenedor Docker y "
  "CloudBeaver, la versión web oficial de DBeaver, para ejecutar los scripts y tomar las capturas. Los datos de "
  "Spring Valley son ficticios y fueron creados con fines académicos.")

h1("Contexto y problemática de Spring Valley")
p("Spring Valley es una empresa (ficticia para efectos del caso) que comercializa bebidas y snacks saludables —aguas, "
  "jugos naturales, tés, bebidas funcionales y snacks— a través de cuatro sucursales en Bogotá, Medellín, Cali y "
  "Barranquilla, con ocho vendedores. El aprendiz fue contratado para guiar a la empresa en la consolidación de sus "
  "datos de ventas, de modo que pueda tomar decisiones con claridad, analizar su historia, hacer pronósticos e "
  "integrar procesos.")
p("El problema central es que todas las consultas se hacían sobre el sistema OLTP (procesamiento de transacciones en "
  "línea). Un sistema OLTP está diseñado para registrar rápidamente operaciones diarias, como una factura, con tablas "
  "normalizadas y muchas escrituras pequeñas. Un sistema OLAP (procesamiento analítico en línea), en cambio, almacena "
  "datos históricos, integrados y desnormalizados para leer y resumir grandes volúmenes sin interrumpir la operación. "
  "Al ejecutar consultas analíticas pesadas sobre el OLTP, la facturación se bloqueó. La bodega de datos separa ambas "
  "cargas de trabajo: el OLTP sigue facturando y la bodega responde las preguntas de la dirección.")
p("Además, la información de ventas llega con problemas de calidad: nombres con mayúsculas mezcladas y espacios, "
  "fechas en distintos formatos, el sexo del cliente y la forma de pago escritos de muchas maneras, y filas "
  "duplicadas. Por eso la bodega necesita un proceso ETL que limpie y estandarice los datos antes de cargarlos.")

h1("Planeación y línea de tiempo")
p("La Tabla 1 presenta la planeación del proyecto. La duración es la propuesta para una implementación real en "
  "Spring Valley, organizada en siete semanas; la columna de estado indica qué se ejecutó en esta evidencia.")
tabla("Línea de tiempo del proyecto de bodega de datos",
      ["Fase", "Actividad", "Semana", "Entregable", "Estado en la evidencia"],
      [["1", "Análisis del problema", "1", "Diagnóstico OLTP vs. OLAP", "Realizado"],
       ["2", "Levantamiento de requerimientos (negocio, compañía y técnicos)", "1", "Requerimientos documentados", "Realizado"],
       ["3", "Identificación de fuentes de datos", "2", "Archivo plano de ventas 2023–2025", "Realizado (datos ficticios)"],
       ["4", "Diseño dimensional (Kimball, grano, hechos y dimensiones)", "2", "Modelo estrella", "Realizado"],
       ["5", "Preparación del entorno (Oracle, Docker, DBeaver)", "3", "Esquema SPRING_VALLEY y conexión", "Realizado"],
       ["6", "Creación del staging", "3", "Tabla STG_VENTAS", "Realizado"],
       ["7", "Proceso ETL: extracción y transformación", "4", "STG_VENTAS_LIMPIA", "Realizado"],
       ["8", "Creación y carga de dimensiones", "4", "Cinco dimensiones cargadas", "Realizado"],
       ["9", "Creación y carga de la tabla de hechos", "5", "FACT_VENTAS cargada", "Realizado"],
       ["10", "Validación de la carga", "5", "Controles de calidad e integridad", "Realizado"],
       ["11", "Consultas analíticas e interpretación", "6", "Respuestas a preguntas de negocio", "Realizado"],
       ["12", "Documentación y entrega", "7", "Informe en PDF", "Realizado"]],
      [1.0, 6.2, 1.5, 4.3, 3.5])

h1("Requerimientos")
h2("Requerimientos del negocio")
p("Son las preguntas que la dirección de Spring Valley necesita responder. Cada una se convirtió en una consulta SQL "
  "sobre la bodega (sección de consultas).")
tabla("Requerimientos del negocio",
      ["Código", "Pregunta del negocio", "Dimensiones y métricas"],
      [["RN1", "¿Cómo evolucionan las ventas por año?", "DIM_TIEMPO; valor_total, cantidad"],
       ["RN2", "¿Qué productos y categorías venden más?", "DIM_PRODUCTO; valor_total, cantidad"],
       ["RN3", "¿Cuánto vende cada sucursal y qué participación tiene?", "DIM_SUCURSAL; valor_total"],
       ["RN4", "¿Qué vendedores tienen mejor desempeño?", "DIM_VENDEDOR; valor_total, facturas"],
       ["RN5", "¿Quiénes son los clientes que más compran y cómo se comportan?", "DIM_CLIENTE; valor_total, ticket"],
       ["RN6", "¿Cómo se distribuyen las ventas por forma de pago?", "forma_pago; valor_total"]],
      [2.0, 9.0, 5.5])
h2("Requerimientos de la compañía (gobierno de datos)")
li("Nombres estandarizados: tablas con prefijos STG_, DIM_ y FACT_; columnas en minúscula con guion bajo.",
   "Textos uniformes: nombres de personas, sucursales y ciudades en formato Nombre Propio (INITCAP); códigos y "
   "productos en mayúsculas (UPPER) y sin espacios sobrantes (TRIM).",
   "Fechas uniformes: todas con tipo DATE, sin importar el formato de origen.",
   "Catálogos controlados: sexo solo M o F; forma de pago solo Efectivo, Tarjeta crédito, Tarjeta débito o Transferencia.",
   "Información sin duplicados: una fila por línea de factura y una fila por cliente, producto, vendedor, sucursal y fecha.",
   "Integridad de datos: llaves primarias, llaves foráneas y restricciones CHECK en Oracle.",
   "Procesos medibles: conteos y totales comparados entre cada etapa del ETL.")
h2("Requerimientos técnicos")
tabla("Herramientas de hardware y software utilizadas",
      ["Recurso", "Detalle", "Uso en el proyecto"],
      [["Entorno", "Máquina virtual cloud, Ubuntu 24.04.4 LTS, x86_64", "Servidor del laboratorio (el PC personal no tenía recursos)"],
       ["CPU", "4 vCPU Intel Xeon 2.80 GHz", "Procesamiento de la base de datos"],
       ["Memoria", "15 GiB de RAM", "Oracle usa ≈ 1,5 GB de SGA; queda margen para las herramientas"],
       ["Almacenamiento", "21 GB libres en disco", "Imágenes de contenedores y datos"],
       ["Contenedores", "Docker 29.6.2", "Aislar Oracle y CloudBeaver, como se indicó en las sesiones"],
       ["Motor de base de datos", "Oracle AI Database 26ai Free 23.26.3 (imagen gvenzl/oracle-free)", "Bodega de datos, PDB FREEPDB1, puerto 1521"],
       ["Cliente gráfico", "CloudBeaver Community 26.2.1 (DBeaver web) en Chromium 141", "Conexión, scripts SQL y capturas"],
       ["Cliente complementario", "Oracle SQL Developer 24.3.1", "Diagrama del modelo estrella (pestaña Model)"],
       ["Carga de archivos", "Oracle SQLcl 26.3 (comando LOAD)", "Extracción del archivo plano al staging"],
       ["Lenguaje", "SQL de Oracle", "DDL, ETL, validación y consultas"],
       ["Tiempo de respuesta esperado", "Menos de 1 segundo por consulta analítica", "Medido en CloudBeaver: entre 22 y 341 ms"]],
      [3.5, 6.5, 6.5])

h1("Metodología Kimball")
p("Se aplicó el enfoque de Ralph Kimball, de abajo hacia arriba (bottom-up): la bodega se construye a partir de las "
  "decisiones reales que el negocio necesita tomar, empezando por un proceso de negocio concreto, y no a partir de la "
  "tecnología. El proceso tiene cuatro pasos.")
tabla("Aplicación de los cuatro pasos de Kimball a Spring Valley",
      ["Paso", "Decisión tomada", "Justificación"],
      [["1. Seleccionar el proceso de negocio", "VENTAS", "Es el área principal del caso y la fuente de las preguntas de la dirección."],
       ["2. Declarar el grano", "Una fila de FACT_VENTAS = un producto vendido en una factura (una línea de factura).", "Es el nivel más detallado disponible; permite sumar por factura, producto, cliente, vendedor, sucursal o fecha."],
       ["3. Identificar las dimensiones", "Tiempo, cliente, producto, vendedor y sucursal.", "Responden ¿cuándo?, ¿quién compra?, ¿qué?, ¿quién vende? y ¿dónde?"],
       ["4. Identificar los hechos", "cantidad, precio_unitario y valor_total.", "Son medidas numéricas que se suman o promedian."]],
      [4.2, 5.8, 6.5])
p("El diseño se desarrolló en los tres niveles explicados: conceptual (las cinco preguntas del negocio y sus "
  "dimensiones, comprensible para la gerencia), lógico (llaves, tipos de datos y relaciones de la Figura 1) y físico "
  "(sentencias CREATE TABLE ejecutadas en Oracle, Figuras 13 a 18).")

h1("Identificación de hechos y dimensiones")
tabla("Hechos (métricas) de FACT_VENTAS",
      ["Hecho", "Tipo Oracle", "Aditividad", "Descripción"],
      [["cantidad", "NUMBER(6)", "Aditiva", "Unidades vendidas del producto en la línea"],
       ["precio_unitario", "NUMBER(12,2)", "No aditiva (se promedia)", "Precio de venta de cada unidad"],
       ["valor_total", "NUMBER(14,2)", "Aditiva", "cantidad × precio_unitario (validado con CHECK)"]],
      [3.2, 3.0, 3.6, 6.7])
p("Además, la tabla de hechos guarda numero_factura como dimensión degenerada (un atributo de la transacción que no "
  "necesita tabla propia) y forma_pago como atributo controlado por una restricción CHECK. La forma de pago es un "
  "catálogo de solo cuatro valores, por lo que no se creó una dimensión aparte; así se respeta el modelo de cinco "
  "dimensiones visto en las sesiones.")
tabla("Dimensiones y la pregunta que responden",
      ["Dimensión", "Pregunta", "Atributos principales"],
      [["DIM_TIEMPO", "¿Cuándo?", "fecha, año, trimestre, mes, nombre del mes, día, día de la semana"],
       ["DIM_CLIENTE", "¿Quién compra?", "código, nombre, sexo, ciudad"],
       ["DIM_PRODUCTO", "¿Qué se vende?", "código, nombre, categoría, precio de lista"],
       ["DIM_VENDEDOR", "¿Quién vende?", "código, nombre"],
       ["DIM_SUCURSAL", "¿Dónde?", "código, nombre, ciudad"]],
      [3.5, 3.0, 10.0])

h1("Diseño del modelo estrella")
p("El modelo estrella ubica la tabla de hechos en el centro y la rodea de las dimensiones. Cada dimensión tiene una "
  "llave primaria sustituta (numérica) y un código de negocio único; la tabla de hechos tiene una llave foránea hacia "
  "cada dimensión. La Figura 1 muestra el diseño lógico.")
fig("fig01_modelo_estrella_diseno.png", "Diseño lógico del modelo estrella de ventas de Spring Valley",
    "Diagrama elaborado en notación Mermaid a partir del diseño. Se observa FACT_VENTAS en el centro con sus "
    "métricas (cantidad, precio_unitario y valor_total), la dimensión degenerada numero_factura y cinco llaves "
    "foráneas hacia DIM_TIEMPO, DIM_CLIENTE, DIM_PRODUCTO, DIM_VENDEDOR y DIM_SUCURSAL. Las relaciones son de uno a "
    "muchos: una fecha, un cliente o un producto aparecen en muchas ventas. Este diseño es la base de la "
    "implementación física.")
tabla("Definición de las tablas: columnas, tipos, llaves y restricciones",
      ["Tabla", "Columna", "Tipo Oracle", "Llave / restricción"],
      [["DIM_TIEMPO", "id_tiempo", "NUMBER(8)", "PK (formato AAAAMMDD)"],
       ["DIM_TIEMPO", "fecha", "DATE", "NOT NULL, UNIQUE"],
       ["DIM_TIEMPO", "anio, trimestre, mes, dia", "NUMBER", "NOT NULL; CHECK trimestre 1–4 y mes 1–12"],
       ["DIM_TIEMPO", "nombre_mes, dia_semana", "VARCHAR2(15)", "NOT NULL"],
       ["DIM_CLIENTE", "id_cliente", "NUMBER(6) IDENTITY", "PK"],
       ["DIM_CLIENTE", "codigo_cliente", "VARCHAR2(10)", "NOT NULL, UNIQUE"],
       ["DIM_CLIENTE", "nombre_cliente, ciudad", "VARCHAR2(100) / (50)", "NOT NULL"],
       ["DIM_CLIENTE", "sexo", "CHAR(1)", "CHECK IN ('M','F')"],
       ["DIM_PRODUCTO", "id_producto", "NUMBER(6) IDENTITY", "PK"],
       ["DIM_PRODUCTO", "codigo_producto", "VARCHAR2(10)", "NOT NULL, UNIQUE"],
       ["DIM_PRODUCTO", "nombre_producto, categoria", "VARCHAR2(100) / (50)", "NOT NULL"],
       ["DIM_PRODUCTO", "precio_lista", "NUMBER(12,2)", "CHECK > 0"],
       ["DIM_VENDEDOR", "id_vendedor", "NUMBER(6) IDENTITY", "PK"],
       ["DIM_VENDEDOR", "codigo_vendedor, nombre_vendedor", "VARCHAR2(10) / (100)", "NOT NULL; código UNIQUE"],
       ["DIM_SUCURSAL", "id_sucursal", "NUMBER(4) IDENTITY", "PK"],
       ["DIM_SUCURSAL", "codigo_sucursal, nombre_sucursal, ciudad", "VARCHAR2", "NOT NULL; código UNIQUE"],
       ["FACT_VENTAS", "id_venta", "NUMBER(10) IDENTITY", "PK"],
       ["FACT_VENTAS", "numero_factura", "VARCHAR2(15)", "NOT NULL; UNIQUE con id_producto"],
       ["FACT_VENTAS", "id_tiempo, id_cliente, id_producto, id_vendedor, id_sucursal", "NUMBER", "FK hacia cada dimensión"],
       ["FACT_VENTAS", "forma_pago", "VARCHAR2(20)", "CHECK de 4 valores"],
       ["FACT_VENTAS", "cantidad / precio_unitario", "NUMBER(6) / NUMBER(12,2)", "CHECK > 0"],
       ["FACT_VENTAS", "valor_total", "NUMBER(14,2)", "CHECK = cantidad × precio_unitario"]],
      [3.2, 5.6, 3.8, 3.9])

h1("Herramientas de hardware y software")
p("Las herramientas se resumieron en la Tabla 3. Se siguió el stack presentado en las sesiones: motor Oracle, "
  "contenedores Docker y DBeaver. Dos decisiones dependieron del entorno cloud y se documentan con transparencia:")
li("DBeaver de escritorio no se pudo descargar: los sitios dbeaver.io, dbeaver.com, GitHub y el repositorio PPA "
   "respondieron con un bloqueo de red (HTTP 403) del entorno. Por eso se usó CloudBeaver Community, la edición "
   "web oficial del mismo proyecto DBeaver, disponible como imagen de Docker (dbeaver/cloudbeaver). Su editor SQL, "
   "su navegador de objetos y su conexión JDBC a Oracle son los de DBeaver.",
   "CloudBeaver no incluye el editor de diagramas ER de DBeaver de escritorio. Para la evidencia gráfica de las "
   "relaciones se usó la pestaña Model de Oracle SQL Developer, que ya estaba instalado (Figura 20).",
   "Para cargar el archivo plano en el staging se usó el comando LOAD de Oracle SQLcl, porque la importación de "
   "archivos de CloudBeaver Community requiere subirlos desde el equipo cliente.")
fig("fig02_conexion_cloudbeaver.png", "Conexión a Oracle desde DBeaver (CloudBeaver)",
    "Se creó la conexión SPRING_VALLEY_DWH con el driver Oracle de DBeaver, la URL JDBC "
    "jdbc:oracle:thin:@//localhost:1521/FREEPDB1 y el usuario spring_valley. Al presionar TEST aparece el mensaje "
    "«Connection is established», con el driver Oracle JDBC 23.2 y el servidor Oracle AI Database 26ai Free 23.26.3. "
    "Esta conexión es el punto de partida de toda la implementación de la bodega.")
fig("fig03_error_ora12505.png", "Error real en el primer intento de conexión (ORA-12505)",
    "En el primer intento se usó el modo de configuración Manual, que interpretó «FREEPDB1» como SID. Oracle "
    "respondió «ORA-12505: … SID FREEPDB1 is not registered with the listener», porque FREEPDB1 es un nombre de "
    "servicio de una base de datos conectable y no un SID. La solución fue cambiar al modo URL con la sintaxis de "
    "servicio //host:puerto/servicio, como muestra la Figura 2.")

h1("Preparación de datos")
p("Como la guía indica que la información debe ser creada por el aprendiz, se generó un archivo plano de ventas "
  "2023–2025 que simula una exportación del sistema OLTP de Spring Valley. Tiene 520 filas (255 facturas, de 1 a 3 "
  "productos por factura) y 17 columnas, con 20 clientes, 12 productos, 8 vendedores y 4 sucursales ficticios. Para "
  "practicar la limpieza enseñada en las sesiones, el archivo contiene intencionalmente:")
li("espacios adicionales al inicio y al final (286 filas con espacios en el nombre o la ciudad del cliente);",
   "mayúsculas y minúsculas mezcladas en nombres, ciudades, categorías y códigos (por ejemplo «c019» y «C019»);",
   "fechas en tres formatos: 2023-01-13, 06/01/2023 y 02-01-2023;",
   "el sexo escrito de 10 formas distintas (M, m, Masculino, MASCULINO, hombre, F, f, Femenino, FEMENINO, fem);",
   "la forma de pago escrita de 16 formas distintas (EFECTIVO, Efect., TC, T. Credito, transf, PSE, etc.);",
   "10 filas duplicadas (la misma línea exportada dos veces).")
p("El archivo original se conservó sin modificaciones (datos/ventas_crudas_spring_valley.csv, SHA-256 1a02eace…).")
fig("fig04_datos_crudos.png", "Datos crudos de ventas de Spring Valley",
    "Se observa el archivo plano original abierto en el navegador. Ya en las primeras filas aparecen los problemas: "
    "la factura FV-10001 repetida, códigos en minúscula (c019, p001), el sexo como FEMENINO, fem, hombre o "
    "MASCULINO, ciudades como «   medellín », fechas como 02-01-2023 y 2023-01-13, y formas de pago como TC, PSE o "
    "Efect. Estos datos no pueden cargarse así en la bodega, porque producirían categorías y conteos falsos.")

h1("Tabla staging")
p("La tabla de staging es un área de trabajo que recibe los datos tal como llegan, antes de limpiarlos. Todas sus "
  "columnas son VARCHAR2, para que ningún valor sea rechazado por su formato y no se pierda información en la "
  "extracción.")
fig("fig05_creacion_stg_ventas.png", "Creación de la tabla STG_VENTAS en DBeaver",
    "En el editor SQL de CloudBeaver se ejecutó CREATE TABLE stg_ventas con las 17 columnas del archivo, todas de "
    "tipo VARCHAR2. El panel Statistics muestra «Executed queries: 1 / 1» sin errores. Usar solo texto en el staging "
    "garantiza que la fecha 06/01/2023 o la cantidad «12» lleguen completas y que la conversión de tipos se haga "
    "después, de forma controlada, en la transformación.")

h1("Proceso ETL")
h2("Extracción")
p("La extracción llevó las 520 filas del archivo plano a STG_VENTAS sin alterar ningún valor, con el comando LOAD "
  "de SQLcl (formato CSV, lotes de 100 filas y confirmación automática).")
code("set loadformat csv\nload stg_ventas datos/ventas_crudas_spring_valley.csv")
fig("fig06_extraccion_load.png", "Extracción del archivo crudo hacia STG_VENTAS",
    "Imagen generada a partir del registro real de la terminal (logs/03_extraccion_staging.log). SQLcl informa "
    "«Number of rows processed: 520», «Number of rows in error: 0» y «SUCCESS: Processed without errors», y la "
    "consulta posterior confirma 520 filas en staging. El conteo coincide con el archivo de origen, lo que demuestra "
    "una extracción completa.")
fig("fig07_datos_staging.png", "Datos cargados en staging antes de limpiar",
    "En DBeaver se consultó STG_VENTAS rodeando los textos con corchetes para hacer visibles los espacios. Se "
    "observan valores como «[ RICARDO JOSÉ PEÑA ]», «[ C008 ]» y «[ Cali ]», el sexo como hombre, m o fem, fechas "
    "en formatos mezclados y formas de pago como Efect., TC o PSE. Esto confirma que el staging conservó el dato "
    "original, condición necesaria para poder auditar la limpieza.")
h2("Transformación")
p("La transformación creó la tabla STG_VENTAS_LIMPIA aplicando las técnicas vistas en las sesiones:")
tabla("Técnicas de limpieza aplicadas",
      ["Problema", "Técnica SQL", "Ejemplo"],
      [["Espacios sobrantes", "TRIM", "«   medellín » → «medellín»"],
       ["Mayúsculas en nombres y ciudades", "INITCAP(TRIM(...))", "«óscar iván londoño» → «Óscar Iván Londoño»"],
       ["Códigos y productos", "UPPER(TRIM(...))", "« c016 » → «C016»"],
       ["Sexo en varias formas", "CASE … IN (...)", "MASCULINO, hombre, m → M"],
       ["Forma de pago en varias formas", "CASE … LIKE", "TC, T. Credito, tarjeta credito → Tarjeta crédito"],
       ["Fechas en tres formatos", "CASE + REGEXP_LIKE + TO_DATE", "06/01/2023 y 2023-01-06 → DATE 06/01/2023"],
       ["Números como texto", "TO_NUMBER", "«12» → 12"],
       ["Filas duplicadas", "SELECT DISTINCT", "520 filas → 510 filas"]],
      [4.5, 5.0, 7.0])
code("""CASE
  WHEN REGEXP_LIKE(TRIM(fecha_factura), '^\\d{4}-\\d{2}-\\d{2}$') THEN TO_DATE(TRIM(fecha_factura), 'YYYY-MM-DD')
  WHEN REGEXP_LIKE(TRIM(fecha_factura), '^\\d{2}/\\d{2}/\\d{4}$') THEN TO_DATE(TRIM(fecha_factura), 'DD/MM/YYYY')
  WHEN REGEXP_LIKE(TRIM(fecha_factura), '^\\d{2}-\\d{2}-\\d{4}$') THEN TO_DATE(TRIM(fecha_factura), 'DD-MM-YYYY')
END AS fecha_factura""")
fig("fig08_transformacion.png", "Ejecución de la transformación en DBeaver",
    "Se ejecutó CREATE TABLE stg_ventas_limpia AS SELECT DISTINCT … con TRIM, UPPER, INITCAP, CASE, REGEXP_LIKE y "
    "TO_DATE. El panel Statistics indica «Updated Rows: 510»: de las 520 filas crudas quedaron 510, porque DISTINCT "
    "eliminó los 10 duplicados. Este paso convierte datos inconsistentes en información confiable antes de tocar la "
    "bodega.")
fig("fig09_limpieza_catalogos.png", "Antes y después: sexo y forma de pago",
    "La consulta une STG_VENTAS con STG_VENTAS_LIMPIA y agrupa cada valor crudo con su valor estandarizado. Las 10 "
    "variantes de sexo ([FEMENINO], [f], [fem ], [ hombre], [MASCULINO], [m]…) quedaron reducidas a F y M, y las 16 "
    "variantes de forma de pago a cuatro valores (Efectivo, Tarjeta crédito, Tarjeta débito y Transferencia). Sin "
    "este paso, el análisis por forma de pago mostraría 16 categorías en lugar de cuatro.")
fig("fig10_limpieza_textos_fechas.png", "Antes y después: nombres, códigos, ciudades y fechas",
    "Para las facturas FV-10001 a FV-10008 se comparan los valores crudos (entre corchetes) con los limpios. "
    "[09/02/2023] y [2023-02-09] se convirtieron en la misma fecha 2023-02-09; [ C016 ] y [c016] en C016; "
    "[ÓSCAR IVÁN LONDOÑO] y [óscar iván londoño] en Óscar Iván Londoño; y [ BOGOTÁ] en Bogotá. Esto garantiza que "
    "un mismo cliente o una misma fecha se registren una sola vez en las dimensiones.")
h2("Carga")
p("Se respetó el orden indicado por los instructores: primero las dimensiones y después la tabla de hechos, porque "
  "cada venta debe apuntar a llaves de dimensión que ya existan.")
fig("fig11_carga_dimensiones.png", "Carga de las dimensiones (primero)",
    "Con INSERT INTO … SELECT DISTINCT desde STG_VENTAS_LIMPIA se poblaron las cinco dimensiones sin duplicados. "
    "DIM_TIEMPO calcula año, trimestre, mes, nombre del mes y día de la semana en español con EXTRACT y TO_CHAR. El "
    "resultado es «Executed queries: 6 / 6» y «Updated Rows: 271» (227 fechas, 20 clientes, 12 productos, "
    "8 vendedores y 4 sucursales), confirmados con COMMIT.")
fig("fig12_carga_hechos.png", "Carga de la tabla de hechos (después)",
    "La carga de FACT_VENTAS resuelve cada llave foránea con un JOIN entre el staging limpio y su dimensión (por "
    "fecha y por código de cliente, producto, vendedor y sucursal) y calcula valor_total = cantidad × "
    "precio_unitario. Se insertaron 510 filas («Updated Rows: 510»), una por cada línea de factura limpia, lo que "
    "cumple el grano declarado.")

h1("Implementación del modelo")
p("Las seis tablas del modelo estrella se crearon en DBeaver antes de la carga. Las figuras siguientes muestran cada "
  "sentencia ejecutada.")
fig("fig13_creacion_dim_tiempo.png", "Creación de DIM_TIEMPO",
    "CREATE TABLE dim_tiempo define id_tiempo NUMBER(8) como llave primaria en formato AAAAMMDD, la fecha única y "
    "los atributos de calendario, con restricciones CHECK para trimestre (1 a 4) y mes (1 a 12). Esta dimensión "
    "permite analizar las ventas por año, trimestre y mes.")
fig("fig14_creacion_dim_cliente.png", "Creación de DIM_CLIENTE",
    "La dimensión de clientes usa una llave sustituta id_cliente generada automáticamente (IDENTITY), el código de "
    "negocio como UNIQUE y la restricción CHECK sexo IN ('M','F'), que obliga a que solo entren valores ya "
    "estandarizados. Responde a la pregunta ¿quién compra?")
fig("fig15_creacion_dim_producto.png", "Creación de DIM_PRODUCTO",
    "Se define el producto con código único, nombre, categoría y precio de lista NUMBER(12,2) con CHECK > 0. La "
    "categoría permite agrupar las ventas por línea de negocio (aguas, jugos, tés, bebidas funcionales y snacks).")
fig("fig16_creacion_dim_vendedor.png", "Creación de DIM_VENDEDOR",
    "La dimensión de vendedores contiene la llave sustituta, el código único y el nombre. Es la base del ranking de "
    "desempeño comercial.")
fig("fig17_creacion_dim_sucursal.png", "Creación de DIM_SUCURSAL",
    "La dimensión de sucursales guarda código, nombre y ciudad, y responde a la pregunta ¿dónde se vende? En todas "
    "las sentencias el panel Statistics muestra «Executed queries: 1 / 1» sin errores.")
fig("fig18_creacion_fact_ventas.png", "Creación de FACT_VENTAS",
    "La tabla de hechos contiene id_venta como llave primaria, numero_factura como dimensión degenerada, cinco "
    "llaves foráneas (FK_FACT_TIEMPO, FK_FACT_CLIENTE, FK_FACT_PRODUCTO, FK_FACT_VENDEDOR y FK_FACT_SUCURSAL) y las "
    "métricas. Las restricciones CHECK validan el catálogo de forma de pago, que cantidad y precio sean positivos y "
    "que valor_total sea exactamente cantidad × precio_unitario; la restricción UNIQUE (numero_factura, "
    "id_producto) impide duplicar una línea de factura.")
fig("fig19_tablas_esquema.png", "Tablas del esquema SPRING_VALLEY en DBeaver",
    "En el navegador de objetos de CloudBeaver, la pestaña Tables del esquema SPRING_VALLEY lista las ocho tablas: "
    "las cinco dimensiones, FACT_VENTAS y las dos tablas de staging (STG_VENTAS y STG_VENTAS_LIMPIA), todas en el "
    "tablespace USERS. Esto confirma la estructura física completa de la bodega.")
fig("fig20_modelo_estrella_sqldeveloper.png", "Esquema estrella generado desde la base de datos",
    "La pestaña Model de FACT_VENTAS en Oracle SQL Developer lee las llaves foráneas reales de Oracle y dibuja la "
    "estrella: FACT_VENTAS (en rojo) se relaciona con DIM_VENDEDOR, DIM_CLIENTE, DIM_PRODUCTO, DIM_SUCURSAL y "
    "DIM_TIEMPO. Se usó SQL Developer porque CloudBeaver Community no trae editor de diagramas ER. El diagrama "
    "coincide con el diseño de la Figura 1, lo que comprueba que el modelo quedó implementado tal como se diseñó.")
fig("fig21_dim_tiempo_cargada.png", "DIM_TIEMPO cargada",
    "La consulta SELECT * FROM dim_tiempo muestra las fechas con su llave AAAAMMDD, el año, el trimestre, el mes, el "
    "nombre del mes y el día de la semana en español (Enero, Lunes, Viernes…). Al tener el calendario precalculado, "
    "las consultas por período no necesitan funciones de fecha.")
fig("fig22_dim_cliente_cargada.png", "DIM_CLIENTE cargada",
    "La dimensión contiene 20 clientes, uno por código, con nombres en formato Nombre Propio, el sexo como M o F y "
    "la ciudad estandarizada. Aunque en el archivo crudo el mismo cliente aparecía escrito de varias formas, aquí "
    "existe una sola fila por cliente.")
fig("fig23_dim_producto_cargada.png", "DIM_PRODUCTO cargada",
    "Se observan los 12 productos (P001 a P012) con su categoría en mayúsculas y su precio de lista, sin espacios ni "
    "variaciones de escritura.")
fig("fig24_dim_vendedor_cargada.png", "DIM_VENDEDOR cargada",
    "La dimensión contiene los 8 vendedores (V01 a V08) con su nombre estandarizado.")
fig("fig25_dim_sucursal_cargada.png", "DIM_SUCURSAL cargada",
    "Se observan las 4 sucursales de Spring Valley: Bogotá Centro, Medellín Laureles, Cali Norte y Barranquilla "
    "Prado, con su ciudad.")
fig("fig26_fact_ventas_cargada.png", "Tabla de hechos FACT_VENTAS cargada",
    "Cada fila de FACT_VENTAS es una línea de factura con sus llaves numéricas hacia las dimensiones, la forma de "
    "pago estandarizada, la cantidad, el precio unitario y el valor total. Los textos descriptivos ya no se repiten "
    "en la tabla de hechos: se obtienen de las dimensiones mediante JOIN.")

h1("Validación de la información")
p("Antes de analizar se validó la carga, como indicaron los instructores. La Tabla 9 resume los controles.")
tabla("Resultados de la validación del ETL",
      ["Control", "Resultado", "Interpretación"],
      [["Registros en staging crudo / limpio", "520 / 510", "Se eliminaron exactamente los 10 duplicados"],
       ["Registros en FACT_VENTAS", "510", "Ninguna línea limpia se perdió en la carga"],
       ["Dimensiones", "227 fechas, 20 clientes, 12 productos, 8 vendedores, 4 sucursales", "Sin duplicados"],
       ["Líneas duplicadas en FACT_VENTAS", "0", "Se cumple el grano"],
       ["Fechas o sexo nulos tras la limpieza", "0", "Los tres formatos de fecha se convirtieron"],
       ["Valores de SEXO", "F, M", "Catálogo válido"],
       ["Valores de FORMA_PAGO", "Efectivo, Transferencia, Tarjeta crédito, Tarjeta débito", "Catálogo válido"],
       ["Rango de fechas", "2023-01-02 a 2025-12-24", "Tres años de historia"],
       ["Ventas huérfanas (sin cliente / sin producto)", "0 / 0", "Integridad referencial completa"],
       ["Total staging limpio vs. FACT_VENTAS", "25.989.200 = 25.989.200", "Métricas coherentes"],
       ["Llaves primarias / foráneas", "6 / 5", "Modelo estrella completo"]],
      [5.5, 5.5, 5.5])
fig("fig27_validacion_conteos.png", "Validación de registros por etapa",
    "Con UNION ALL de varios COUNT(*) se compararon las etapas del ETL: 520 filas crudas, 510 limpias, 510 en "
    "FACT_VENTAS y el número de filas de cada dimensión. La igualdad entre el staging limpio y la tabla de hechos "
    "demuestra que la carga fue completa.")
fig("fig28_validacion_calidad.png", "Validación de calidad: duplicados, nulos, catálogos y métricas",
    "La consulta reúne nueve controles. Muestra 0 líneas duplicadas, 0 fechas o sexos nulos, los catálogos válidos "
    "de sexo y forma de pago, el rango de fechas, 0 ventas huérfanas y el mismo total (25.989.200) en el staging "
    "limpio y en FACT_VENTAS. La pantalla recorta algunos textos largos; el resultado completo está en "
    "logs/04_validacion_y_consultas.log y en la Tabla 9.")
fig("fig29_validacion_llaves.png", "Validación de llaves primarias y foráneas",
    "Desde las vistas USER_CONSTRAINTS y USER_CONS_COLUMNS se listaron las 6 llaves primarias (una por tabla del "
    "modelo) y las 5 llaves foráneas de FACT_VENTAS con su tabla referenciada. Esto confirma que la integridad "
    "referencial está definida en el motor y no depende de la aplicación.")
fig("fig30_prueba_integridad.png", "Prueba de integridad referencial (ORA-02291)",
    "Se intentó registrar una venta para el cliente 9999, que no existe en DIM_CLIENTE. Oracle la rechazó con "
    "«ORA-02291: integrity constraint (SPRING_VALLEY.FK_FACT_CLIENTE) violated - parent key not found». La consulta "
    "posterior confirmó que FACT_VENTAS siguió con 510 filas, de modo que los datos válidos no se alteraron. Esta "
    "prueba demuestra que la bodega no admite ventas sin contexto.")

h1("Consultas y análisis de negocio")
p("Con la bodega validada se respondieron las preguntas del negocio. Todas las consultas unen FACT_VENTAS con sus "
  "dimensiones (JOIN), agrupan (GROUP BY) y calculan métricas (SUM, COUNT, AVG y funciones analíticas).")
code("""SELECT t.anio, COUNT(DISTINCT f.numero_factura) AS facturas, SUM(f.cantidad) AS unidades,
       SUM(f.valor_total) AS ventas, ROUND(AVG(f.valor_total)) AS promedio_por_linea,
       ROUND(100 * (SUM(f.valor_total) / LAG(SUM(f.valor_total)) OVER (ORDER BY t.anio) - 1), 1) AS crecimiento_pct
  FROM fact_ventas f JOIN dim_tiempo t ON t.id_tiempo = f.id_tiempo
 GROUP BY t.anio ORDER BY t.anio;""")
fig("fig31_indicadores_generales.png", "Primera consulta analítica: indicadores generales",
    "La primera consulta resume la bodega: 255 facturas, 510 líneas, 3.638 unidades, ventas por $25.989.200, ticket "
    "promedio de $101.918 por factura y 20 clientes, entre el 2 de enero de 2023 y el 24 de diciembre de 2025. Estos "
    "indicadores son la línea base contra la que la dirección podrá medir resultados futuros.")
fig("fig32_ventas_por_anio.png", "Evolución de las ventas por año",
    "Uniendo FACT_VENTAS con DIM_TIEMPO y agrupando por año, las ventas pasaron de $6.419.400 (2023) a $8.707.400 "
    "(2024) y a $10.862.400 (2025). La función LAG calcula un crecimiento del 35,6 % en 2024 y del 24,7 % en 2025. "
    "El negocio crece de forma sostenida, aunque a menor ritmo; además, el valor promedio por línea subió de $46.857 a "
    "$56.282. Esta tendencia sirve de base para pronosticar 2026 y planear compras e inventario.")
fig("fig33_productos.png", "Productos con mayores ventas",
    "Con RANK() sobre SUM(valor_total), los tres primeros lugares corresponden a snacks saludables: Mix de frutos "
    "secos ($4.801.800), Barra de cereal ($4.337.500) y Galletas integrales ($2.625.500). Las aguas venden muchas "
    "unidades (Agua con gas, 363) pero poco valor por su bajo precio. Spring Valley debería priorizar el "
    "abastecimiento y la promoción de snacks y jugos, que generan más ingresos por unidad.")
fig("fig34_sucursales.png", "Ventas por sucursal",
    "Medellín Laureles lidera con $7.326.000 (28,2 % de las ventas y 75 facturas), seguida de Cali Norte (25,9 %) y "
    "Barranquilla Prado (25,0 %); Bogotá Centro es la de menor participación (21,0 %) y menor ticket promedio "
    "($95.628). La empresa podría replicar las prácticas de Medellín y revisar la estrategia comercial de Bogotá.")
fig("fig35_vendedores.png", "Ranking de vendedores",
    "El ranking ordena a los vendedores por valor vendido. Fabián Alberto Rincón (Cali Norte) ocupa el primer lugar "
    "con $4.021.000 y el ticket promedio más alto ($129.710), aunque con menos facturas (31) que Claudia Marcela "
    "Henao (39). En la misma sucursal, Gloria Inés Muñoz ocupa el último lugar ($2.700.100), lo que muestra "
    "diferencias de desempeño que pueden orientar capacitaciones e incentivos.")
fig("fig36_clientes.png", "Clientes que más compran",
    "Los 10 clientes con mayor valor comprado están encabezados por Julián David Pardo ($2.208.800 en 12 compras, "
    "ticket de $184.067). Sebastián Álvarez Gil es el que más veces compró (21), pero con un ticket bajo ($68.962). "
    "Esto sugiere dos estrategias: fidelizar a los clientes de alto ticket y aumentar el valor de compra de los "
    "clientes frecuentes.")
fig("fig37_clientes_sexo.png", "Ventas por sexo del cliente",
    "Agrupando por el sexo ya estandarizado, los 10 clientes masculinos concentran el 56,7 % de las ventas y las "
    "10 clientes femeninas el 43,3 %. Este análisis solo es posible porque el ETL convirtió las 10 formas de escribir "
    "el sexo en M y F.")
fig("fig38_forma_pago.png", "Distribución de ventas por forma de pago",
    "La tarjeta de crédito concentra el 38,4 % de las ventas con el ticket más alto ($116.128), aunque el efectivo "
    "tiene más facturas (91 frente a 86). Tarjeta débito (15,3 %) y transferencia (14,2 %) completan la "
    "distribución. Spring Valley puede usar este dato para negociar comisiones con los bancos y promover los medios "
    "electrónicos.")

h1("Resultados obtenidos")
tabla("Síntesis de resultados",
      ["Aspecto", "Resultado"],
      [["Bodega implementada", "Esquema SPRING_VALLEY en Oracle con 1 tabla de hechos, 5 dimensiones y 2 tablas de staging"],
       ["ETL", "520 filas crudas → 510 filas limpias → 510 hechos; 10 duplicados eliminados"],
       ["Calidad", "Sexo: de 10 variantes a 2; forma de pago: de 16 variantes a 4; fechas: de 3 formatos a DATE"],
       ["Integridad", "6 PK, 5 FK, CHECK y UNIQUE; 0 huérfanos; ORA-02291 ante una venta inválida"],
       ["Ventas 2023–2025", "$25.989.200; crecimiento del 35,6 % en 2024 y del 24,7 % en 2025"],
       ["Hallazgos", "Snacks saludables lideran; Medellín es la sucursal principal; la tarjeta de crédito es la forma de pago de mayor valor"],
       ["Rendimiento", "Consultas analíticas resueltas entre 22 y 341 ms, sin tocar el sistema transaccional"]],
      [4.5, 12.0])

h1("Conclusiones")
p("La construcción de la bodega de datos de Spring Valley demostró que separar el análisis del sistema transaccional "
  "resuelve el problema que originó el caso: las consultas históricas se ejecutan sobre un esquema propio, en "
  "milisegundos, sin competir con la facturación.")
p("La metodología Kimball permitió pasar de las preguntas de la dirección a un diseño concreto. Declarar el grano "
  "(una línea de factura) fue la decisión más importante, porque definió qué representa cada fila y garantizó que "
  "las sumas por cualquier dimensión fueran correctas. El modelo estrella, con FACT_VENTAS y cinco dimensiones, "
  "resultó simple de consultar: cada pregunta se respondió con unos pocos JOIN y un GROUP BY.")
p("El proceso ETL fue la etapa más exigente y la de mayor valor. Los datos crudos tenían 10 formas de escribir el "
  "sexo, 16 de escribir la forma de pago, tres formatos de fecha y filas duplicadas; con TRIM, UPPER, INITCAP, CASE y "
  "TO_DATE se estandarizaron por completo. Sin esa limpieza, los análisis habrían mostrado categorías falsas y "
  "totales inflados.")
p("La validación y las restricciones de Oracle dieron confianza en los resultados: los conteos y totales coinciden "
  "entre etapas, no hay claves huérfanas y el motor rechazó una venta con un cliente inexistente. Finalmente, las "
  "consultas analíticas entregaron información útil para decidir: crecimiento sostenido, liderazgo de los snacks "
  "saludables y de la sucursal de Medellín, diferencias entre vendedores y preferencia por la tarjeta de crédito.")

h1("Oportunidades de mejora")
li("Automatizar el ETL con procedimientos PL/SQL y tareas programadas (DBMS_SCHEDULER) para no depender de "
   "ejecuciones manuales.",
   "Actualizar la bodega periódicamente (por ejemplo, cada noche) con cargas incrementales que solo traigan las "
   "ventas nuevas.",
   "Incorporar controles de calidad automáticos que detengan la carga si aparece un valor fuera de catálogo o una "
   "fecha inválida, y registrar las filas rechazadas en una tabla de errores.",
   "Aplicar dimensiones de cambio lento (SCD tipo 2) para conservar la historia cuando un cliente cambie de ciudad o "
   "un vendedor de sucursal.",
   "Conectar herramientas de visualización (Power BI, Tableau u Oracle Analytics) para construir tableros para la "
   "gerencia.",
   "Definir políticas de respaldo (RMAN o Data Pump) y de recuperación de la bodega.",
   "Reforzar la seguridad: usuarios de solo lectura para analistas, contraseñas robustas y cifrado de conexiones.",
   "Planear el crecimiento: agregar nuevos procesos de negocio (inventario, compras) que compartan las dimensiones "
   "conformadas de tiempo, producto y sucursal, y particionar FACT_VENTAS por año cuando aumente el volumen.")

h1("Referencias")
for r in [
    "Arévalo Vela, M. A. (2026a, 22 de septiembre). <i>Resumen sesión en línea fase 1 análisis: Explicación evidencia técnica GA4-210602033-AA1-EV03</i> [Documento de sesión]. Servicio Nacional de Aprendizaje (SENA).",
    "Arévalo Vela, M. A. (2026b, 23 de septiembre). <i>Resumen sesión en línea fase 1 análisis: Explicación evidencia técnica GA4-210602033-AA1-EV03, parte 2</i> [Documento de sesión]. Servicio Nacional de Aprendizaje (SENA).",
    "DBeaver Corporation. (s.f.). <i>CloudBeaver documentation</i>. https://dbeaver.com/docs/cloudbeaver/",
    "Fuel Tulcán, C. A. (2026a, 24 de septiembre). <i>Resumen sesión en línea: Explicación evidencia GA4-210602033-AA1-EV03, parte 3</i> [Documento de sesión]. Servicio Nacional de Aprendizaje (SENA).",
    "Fuel Tulcán, C. A. (2026b, 25 de septiembre). <i>Resumen sesión en línea: Explicación evidencia GA4-210602033-AA1-EV03, parte 4</i> [Documento de sesión]. Servicio Nacional de Aprendizaje (SENA).",
    "Inmon, W. H. (2005). <i>Building the data warehouse</i> (4.ª ed.). Wiley.",
    "Kimball, R., y Ross, M. (2013). <i>The data warehouse toolkit: The definitive guide to dimensional modeling</i> (3.ª ed.). Wiley.",
    "Oracle. (s.f.). <i>Oracle Database SQL language reference</i>. https://docs.oracle.com/en/database/oracle/oracle-database/23/sqlrf/",
    "Oracle. (s.f.). <i>Oracle SQLcl user's guide</i>. https://docs.oracle.com/en/database/oracle/sql-developer-command-line/",
    "Servicio Nacional de Aprendizaje (SENA). (s.f.). <i>Bodegas de datos</i> [Componente formativo]. Programa Implementación y gestión de bases de datos.",
]:
    B.append(("ref", r))

h1("Anexo. Lista de chequeo de la evidencia")
tabla("Control final de requisitos",
      ["Requisito", "Cumple", "Evidencia / sección"],
      [["Caso Spring Valley, área de ventas", "Sí", "Contexto; todo el documento"],
       ["Planeación con línea de tiempo", "Sí", "Planeación, Tabla 1"],
       ["Herramientas de hardware y software", "Sí", "Requerimientos técnicos, Tabla 3; Herramientas"],
       ["Requerimientos de negocio, compañía y técnicos", "Sí", "Requerimientos, Tablas 2 y 3"],
       ["Información relevante para ventas", "Sí", "Requerimientos del negocio; Hechos y dimensiones"],
       ["Tablas, columnas, tipos de datos y relaciones", "Sí", "Diseño del modelo, Tabla 7; Figuras 13 a 18"],
       ["Esquema estrella", "Sí", "Figuras 1 y 20"],
       ["Staging y proceso ETL", "Sí", "Figuras 4 a 12"],
       ["Dimensiones primero, hechos después", "Sí", "Figuras 11 y 12"],
       ["Validación e integridad referencial", "Sí", "Tabla 9; Figuras 27 a 30"],
       ["Capturas reales con interpretación", "Sí", "38 figuras con párrafo explicativo"],
       ["Consultas de negocio interpretadas", "Sí", "Figuras 31 a 38"],
       ["Conclusiones y oportunidades de mejora", "Sí", "Secciones finales"],
       ["Referencias en APA 7", "Sí", "Referencias"],
       ["Mínimo 8 páginas, formato PDF", "Sí", "Archivo PDF generado"]],
      [7.0, 1.8, 7.7])

# ============================ SALIDA WORD ============================
def generar_docx(ruta):
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Inches(8.5), Inches(11)
    for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(s, m, Inches(1))
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"; st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    st.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    st.paragraph_format.space_after = Pt(0)
    hp = s.header.paragraphs[0]; hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = hp.add_run()
    for tag, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        el = OxmlElement("w:fldChar" if tag else "w:instrText")
        if tag: el.set(qn("w:fldCharType"), tag)
        else: el.set(qn("xml:space"), "preserve"); el.text = txt
        r._r.append(el)

    def centrado(t, bold=False, espacio=0):
        q = doc.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.CENTER
        q.paragraph_format.space_before = Pt(espacio)
        q.add_run(t).bold = bold
    centrado("Realización de caso estudio bodega de datos: Spring Valley", True, 120)
    centrado("Evidencia GA4-210602033-AA1-EV03")
    for t in ["[Nombre del aprendiz]", "Servicio Nacional de Aprendizaje (SENA)",
              "Tecnólogo en Implementación y gestión de bases de datos — Ficha [número de ficha]",
              "Instructores: Miguel Andrés Arévalo Vela y Carlos Alberto Fuel Tulcán",
              "[Centro de formación], [Ciudad]", "[Fecha de entrega]"]:
        centrado(t)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    nfig, ntab = [1], [1]
    for tipo, d in B:
        if tipo == "h1":
            q = doc.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.CENTER
            q.paragraph_format.keep_with_next = True; q.add_run(d).bold = True
        elif tipo == "h2":
            q = doc.add_paragraph(); q.paragraph_format.keep_with_next = True; q.add_run(d).bold = True
        elif tipo == "p":
            q = doc.add_paragraph(); q.paragraph_format.first_line_indent = Inches(0.5); q.add_run(d)
        elif tipo == "ul":
            for it in d:
                q = doc.add_paragraph(style="List Bullet")
                q.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE; q.add_run(it)
        elif tipo == "code":
            q = doc.add_paragraph(); q.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            q.paragraph_format.left_indent = Inches(0.3); q.paragraph_format.space_after = Pt(8)
            rr = q.add_run(d); rr.font.name = "Courier New"; rr.font.size = Pt(8.5)
        elif tipo == "tabla":
            titulo, enc, filas, anchos = d
            q = doc.add_paragraph(); q.paragraph_format.keep_with_next = True; q.add_run(f"Tabla {ntab[0]}").bold = True
            q = doc.add_paragraph(); q.paragraph_format.keep_with_next = True; q.add_run(titulo).italic = True
            ntab[0] += 1
            t = doc.add_table(rows=1, cols=len(enc)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for i, e in enumerate(enc):
                t.rows[0].cells[i].text = ""; t.rows[0].cells[i].paragraphs[0].add_run(e).bold = True
            for f in filas:
                cs = t.add_row().cells
                for i, v in enumerate(f): cs[i].text = str(v)
            for row in t.rows:
                for i, c in enumerate(row.cells):
                    if anchos: c.width = Cm(anchos[i])
                    for par in c.paragraphs:
                        par.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
                        for rr in par.runs: rr.font.size = Pt(9.5)
            doc.add_paragraph()
        elif tipo == "fig":
            arch, titulo, texto = d
            q = doc.add_paragraph(); q.paragraph_format.keep_with_next = True; q.add_run(f"Figura {nfig[0]}").bold = True
            q = doc.add_paragraph(); q.paragraph_format.keep_with_next = True; q.add_run(titulo).italic = True
            nfig[0] += 1
            im = Image.open(os.path.join(CAP, arch)); ancho = 6.3
            if im.height / im.width * ancho > 7.0: ancho = 7.0 * im.width / im.height
            q = doc.add_paragraph(); q.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            q.paragraph_format.keep_with_next = True; q.add_run().add_picture(os.path.join(CAP, arch), width=Inches(ancho))
            q = doc.add_paragraph(); rr = q.add_run("Nota."); rr.italic = True; q.add_run(" Elaboración propia. " + texto)
        elif tipo == "ref":
            q = doc.add_paragraph(); q.paragraph_format.left_indent = Inches(0.5)
            q.paragraph_format.first_line_indent = Inches(-0.5)
            partes = d.split("<i>")
            q.add_run(partes[0])
            for parte in partes[1:]:
                cursiva, resto = parte.split("</i>")
                q.add_run(cursiva).italic = True; q.add_run(resto)
    doc.save(ruta)
    return nfig[0] - 1, ntab[0] - 1


# ============================ SALIDA PDF ============================
CSS = """
@page { size: Letter; margin: 2.54cm; }
body { font-family: 'Liberation Serif', 'Times New Roman', serif; font-size: 12pt; line-height: 2; color: #000; }
h1 { font-size: 12pt; text-align: center; margin: 18px 0 4px; break-after: avoid; }
h2 { font-size: 12pt; margin: 12px 0 2px; break-after: avoid; }
p { text-indent: 1.27cm; margin: 0; text-align: left; }
ul { margin: 0 0 0 1cm; padding: 0; } li { margin: 0; }
.tn, .fn { font-weight: bold; text-indent: 0; margin-top: 10px; break-after: avoid; }
.tt, .ft { font-style: italic; text-indent: 0; break-after: avoid; }
table { border-collapse: collapse; width: 100%; font-size: 9.5pt; line-height: 1.25; margin: 2px 0 12px; }
th, td { border: 1px solid #444; padding: 3px 5px; vertical-align: top; text-align: left; }
th { background: #eee; } tr { break-inside: avoid; }
.img { text-align: center; text-indent: 0; line-height: 1; break-inside: avoid; break-after: avoid; }
.img img { max-width: 100%; max-height: 17.5cm; border: 1px solid #888; }
.nota { text-indent: 0; margin-bottom: 6px; }
pre { font-family: 'DejaVu Sans Mono', monospace; font-size: 7.6pt; line-height: 1.3; background: #f5f5f5;
      border: 1px solid #ccc; padding: 6px 8px; white-space: pre-wrap; margin: 4px 0 10px; }
.ref { text-indent: -1.27cm; padding-left: 1.27cm; }
.portada { text-align: center; break-after: page; padding-top: 5cm; }
.portada p { text-indent: 0; text-align: center; }
"""

def generar_pdf(ruta):
    from playwright.sync_api import sync_playwright
    e = html.escape
    out = ["<div class='portada'><p><b>Realización de caso estudio bodega de datos: Spring Valley</b></p>",
           "<p>Evidencia GA4-210602033-AA1-EV03</p><br>"]
    for t in ["[Nombre del aprendiz]", "Servicio Nacional de Aprendizaje (SENA)",
              "Tecnólogo en Implementación y gestión de bases de datos — Ficha [número de ficha]",
              "Instructores: Miguel Andrés Arévalo Vela y Carlos Alberto Fuel Tulcán",
              "[Centro de formación], [Ciudad]", "[Fecha de entrega]"]:
        out.append(f"<p>{e(t)}</p>")
    out.append("</div>")
    nf, nt = 1, 1
    for tipo, d in B:
        if tipo == "h1": out.append(f"<h1>{e(d)}</h1>")
        elif tipo == "h2": out.append(f"<h2>{e(d)}</h2>")
        elif tipo == "p": out.append(f"<p>{e(d)}</p>")
        elif tipo == "ul": out.append("<ul>" + "".join(f"<li>{e(i)}</li>" for i in d) + "</ul>")
        elif tipo == "code": out.append(f"<pre>{e(d)}</pre>")
        elif tipo == "ref": out.append(f"<p class='ref'>{d}</p>")
        elif tipo == "tabla":
            titulo, enc, filas, _ = d
            out.append(f"<p class='tn'>Tabla {nt}</p><p class='tt'>{e(titulo)}</p><table><tr>"
                       + "".join(f"<th>{e(x)}</th>" for x in enc) + "</tr>"
                       + "".join("<tr>" + "".join(f"<td>{e(str(x))}</td>" for x in f) + "</tr>" for f in filas)
                       + "</table>")
            nt += 1
        elif tipo == "fig":
            arch, titulo, texto = d
            b64 = base64.b64encode(open(os.path.join(CAP, arch), "rb").read()).decode()
            out.append(f"<p class='fn'>Figura {nf}</p><p class='ft'>{e(titulo)}</p>"
                       f"<p class='img'><img src='data:image/png;base64,{b64}'></p>"
                       f"<p class='nota'><i>Nota.</i> Elaboración propia. {e(texto)}</p>")
            nf += 1
    doc = f"<!doctype html><html lang='es'><head><meta charset='utf-8'><style>{CSS}</style></head><body>{''.join(out)}</body></html>"
    tmp = ruta + ".html"
    open(tmp, "w", encoding="utf-8").write(doc)
    with sync_playwright() as pw:
        b = pw.chromium.launch(executable_path=CHROMIUM)
        pg = b.new_page(); pg.goto("file://" + tmp, wait_until="networkidle")
        pg.pdf(path=ruta, format="Letter", display_header_footer=True,
               header_template="<div style='font-size:10pt;width:100%;text-align:right;margin-right:2.54cm;font-family:serif'><span class='pageNumber'></span></div>",
               footer_template="<div></div>",
               margin={"top": "2.54cm", "bottom": "2.54cm", "left": "2.54cm", "right": "2.54cm"})
        b.close()
    os.remove(tmp)


if __name__ == "__main__":
    nf, nt = generar_docx(os.path.join(RAIZ, NOMBRE + ".docx"))
    generar_pdf(os.path.join(RAIZ, NOMBRE + ".pdf"))
    from pypdf import PdfReader
    print(f"Figuras: {nf} | Tablas: {nt} | Páginas PDF: {len(PdfReader(os.path.join(RAIZ, NOMBRE + '.pdf')).pages)}")
