"""
Genera el documento Word del caso de estudio TECH TRANSITION (evidencia GA4-210602033-AA1-EV02)
con el formato APA del informe del aprendiz: Times New Roman 12, doble espacio, "Figura N" en
negrita, titulo en cursiva, imagen y "Nota. Elaboracion propia. ...".

Las figuras son las capturas reales de capturas_caso/ (SQL Developer, esquema TECH_TRANSITION).
La numeracion de figuras inicia en 17 para continuar despues de la Figura 16 del informe.
"""
import os
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt

RAIZ = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
CAP = os.path.join(RAIZ, "capturas_caso")
SALIDA = os.path.join(RAIZ, "Caso_TECH_TRANSITION_EV02.docx")

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, m, Inches(1))

normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(12)
normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
normal.paragraph_format.space_after = Pt(0)
normal.paragraph_format.space_before = Pt(0)

# Numero de pagina arriba a la derecha (como el informe)
hp = sec.header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = hp.add_run()
for tag, txt in (("begin", None), (None, "PAGE"), ("end", None)):
    if tag:
        el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), tag)
    else:
        el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = txt
    run._r.append(el)

fig_num = [17]

# Recortes (x0, y0, x1, y1) sobre la pantalla de 1600x900: se elimina el panel lateral o zonas vacias
# para que el codigo y los resultados se lean bien en la pagina. No se modifica el contenido.
RECORTES = {
    17: (290, 80, 1000, 900), 18: (452, 212, 1232, 712),
    19: (290, 80, 1100, 700), 20: (290, 80, 1100, 700), 21: (290, 80, 1100, 700),
    22: (290, 80, 1100, 700), 23: (290, 80, 1100, 700), 24: (290, 80, 1100, 700),
    25: (290, 80, 1100, 900), 26: (0, 80, 1135, 480),
    40: (290, 80, 1100, 700), 41: (290, 80, 1000, 700), 42: (290, 80, 1600, 850),
    43: (290, 80, 1330, 620), 44: (290, 80, 1000, 520), 45: (290, 80, 1000, 580),
    46: (290, 80, 1000, 480), 47: (290, 80, 1250, 560), 48: (290, 80, 1120, 650),
    49: (290, 80, 1000, 500), 50: (290, 80, 1000, 480), 51: (290, 80, 1000, 480),
    52: (290, 80, 1300, 900),
}
RECORTE_DEFECTO = (290, 80, 1600, 900)
DIR_RECORTES = os.path.join(CAP, "recortes")
os.makedirs(DIR_RECORTES, exist_ok=True)
tab_num = [1]


def h1(texto):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(texto).bold = True


def h2(texto):
    p = doc.add_paragraph()
    p.add_run(texto).bold = True


def parrafo(texto, sangria=True):
    p = doc.add_paragraph()
    if sangria:
        p.paragraph_format.first_line_indent = Inches(0.5)
    p.add_run(texto)
    return p


def vinetas(items):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
        p.add_run(it)


def figura(archivo, titulo, nota):
    n = fig_num[0]; fig_num[0] += 1
    p = doc.add_paragraph(); p.add_run(f"Figura {n}").bold = True
    p.paragraph_format.keep_with_next = True
    p = doc.add_paragraph(); p.add_run(titulo).italic = True
    p.paragraph_format.keep_with_next = True
    p = doc.add_paragraph(); p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.paragraph_format.keep_with_next = True
    from PIL import Image
    img = Image.open(os.path.join(CAP, archivo)).crop(RECORTES.get(n, RECORTE_DEFECTO))
    ruta = os.path.join(DIR_RECORTES, archivo); img.save(ruta)
    ancho = 6.3
    if img.height / img.width * ancho > 7.0:
        ancho = 7.0 * img.width / img.height
    p.add_run().add_picture(ruta, width=Inches(ancho))
    p = doc.add_paragraph()
    r = p.add_run("Nota."); r.italic = True
    p.add_run(" Elaboración propia. " + nota)


def tabla(titulo, encabezados, filas, anchos=None):
    n = tab_num[0]; tab_num[0] += 1
    p = doc.add_paragraph(); p.add_run(f"Tabla {n}").bold = True
    p.paragraph_format.keep_with_next = True
    p = doc.add_paragraph(); p.add_run(titulo).italic = True
    p.paragraph_format.keep_with_next = True
    t = doc.add_table(rows=1, cols=len(encabezados))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, e in enumerate(encabezados):
        c = t.rows[0].cells[i]; c.text = ""
        r = c.paragraphs[0].add_run(e); r.bold = True
    for f in filas:
        cells = t.add_row().cells
        for i, v in enumerate(f):
            cells[i].text = str(v)
    for row in t.rows:
        for i, c in enumerate(row.cells):
            if anchos:
                c.width = Cm(anchos[i])
            for par in c.paragraphs:
                par.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
                for r in par.runs:
                    r.font.size = Pt(10)
    doc.add_paragraph()


# ---------------------------------------------------------------------------
h1("Caso de estudio TECH TRANSITION")
parrafo("TECH TRANSITION es una empresa que comercializa equipos de cómputo, periféricos, elementos de red y "
        "servicios técnicos en tres sucursales ubicadas en Bogotá, Medellín y Cali. La empresa seleccionó al aprendiz "
        "para crear una base de datos para su departamento de facturación, que sirva como punto de partida para una "
        "futura bodega de datos. Según la guía, toda la información del caso debe ser creada por el aprendiz; por eso, "
        "los clientes, documentos, teléfonos, direcciones y correos utilizados son ficticios y se crearon con fines "
        "académicos (los correos usan el dominio reservado example.com).")
parrafo("El caso pide elaborar una lista de chequeo con los requerimientos, verificar si la información está recopilada "
        "y organizada, crear la base de datos en Oracle por medio de SQL Developer y crear las tablas con sus columnas, "
        "tipos de datos y relaciones para almacenar el número de factura, la fecha de factura, la fecha de vencimiento, "
        "el cliente, la sucursal, el valor total, la forma de pago, el valor pagado, el artículo, la cantidad y el valor "
        "unitario. Para su desarrollo se utilizó Oracle Database Free, en la base de datos conectable FREEPDB1, "
        "administrada desde Oracle SQL Developer 24.3.1 en un entorno cloud.")

h2("Lista de chequeo")
parrafo("La lista de chequeo reúne los requerimientos que la base de datos debía solucionar. Se diligenció al terminar "
        "el trabajo, con base en los resultados obtenidos en Oracle.")
tabla("Lista de chequeo de requerimientos del caso TECH TRANSITION",
      ["N.º", "Requerimiento", "Cumple / No cumple", "Observaciones"],
      [["1", "Usuario o esquema del caso creado", "Cumple", "Usuario TECH_TRANSITION con privilegios y cuota de 200 MB (Figura 17)."],
       ["2", "Conexión desde SQL Developer", "Cumple", "La conexión TECH_TRANSITION muestra «Estado: Correcto» (Figura 18)."],
       ["3", "Información de clientes", "Cumple", "Tabla CLIENTES con 10 registros (Figuras 19, 27 y 28)."],
       ["4", "Información de sucursales", "Cumple", "Tabla SUCURSALES con 3 registros (Figuras 20, 29 y 30)."],
       ["5", "Información de formas de pago", "Cumple", "Tabla FORMAS_PAGO con 6 registros (Figuras 21, 31 y 32)."],
       ["6", "Información de artículos", "Cumple", "Tabla ARTICULOS con 15 registros (Figuras 22, 33 y 34)."],
       ["7", "Número de factura, fecha de factura, fecha de vencimiento, cliente, sucursal, forma de pago, valor total y valor pagado", "Cumple", "Tabla FACTURAS con 20 registros (Figuras 23, 35 y 38)."],
       ["8", "Artículo, cantidad y valor unitario", "Cumple", "Tabla DETALLE_FACTURA con 52 registros (Figuras 24, 36 y 39)."],
       ["9", "Llaves primarias", "Cumple", "Una llave primaria por tabla (Figura 41)."],
       ["10", "Relaciones entre tablas", "Cumple", "Cinco llaves foráneas (Figuras 40 y 41)."],
       ["11", "Integridad de los datos", "Cumple", "Restricciones NOT NULL, UNIQUE y CHECK; Oracle rechazó los datos inválidos (Figura 52)."],
       ["12", "Datos de prueba guardados", "Cumple", "106 registros confirmados con COMMIT."],
       ["13", "Consultas de comprobación", "Cumple", "JOIN, totales, saldos y artículos más vendidos (Figuras 42 a 48)."],
       ["14", "Actualización y eliminación de datos", "Cumple", "UPDATE y DELETE ejecutados correctamente (Figuras 49 a 51)."]],
      [1.2, 5.0, 2.6, 7.2])

h2("Verificación de la información y análisis del caso")
parrafo("Al revisar el caso se determinó que la información no estaba recopilada ni organizada para administrarse en "
        "una base de datos. Como escenario de trabajo se planteó que la facturación se llevaba en hojas de cálculo "
        "separadas por sucursal, en las que cada fila repetía los datos del cliente, del artículo y de la forma de pago. "
        "Esta forma de trabajo generaba datos duplicados y errores de digitación, no controlaba que el número de factura "
        "fuera único, obligaba a calcular a mano el valor total y el saldo pendiente, y no impedía registrar datos "
        "inválidos, como cantidades negativas o pagos mayores que el valor de la factura.")
parrafo("Por esa razón, la información se creó y se organizó en tablas relacionadas. Cada dato se guarda una sola vez "
        "y se vincula con los demás por medio de llaves primarias y foráneas, lo que reduce la redundancia y permite que "
        "el propio motor de Oracle valide la información. La Tabla 2 muestra en qué tabla quedó cada dato exigido.")
tabla("Ubicación de los datos exigidos por el caso",
      ["Dato exigido", "Tabla", "Columna"],
      [["Número de factura", "FACTURAS", "numero_factura"],
       ["Fecha de factura", "FACTURAS", "fecha_factura"],
       ["Fecha de vencimiento", "FACTURAS", "fecha_vencimiento"],
       ["Cliente", "CLIENTES (relacionada con FACTURAS)", "id_cliente"],
       ["Sucursal", "SUCURSALES (relacionada con FACTURAS)", "id_sucursal"],
       ["Valor total", "FACTURAS", "valor_total"],
       ["Forma de pago", "FORMAS_PAGO (relacionada con FACTURAS)", "id_forma_pago"],
       ["Valor pagado", "FACTURAS", "valor_pagado"],
       ["Artículo", "ARTICULOS (relacionada con DETALLE_FACTURA)", "id_articulo"],
       ["Cantidad", "DETALLE_FACTURA", "cantidad"],
       ["Valor unitario", "DETALLE_FACTURA", "valor_unitario"]],
      [4.5, 7.5, 4.0])

h2("Diseño de la base de datos")
parrafo("El diseño se normalizó hasta la tercera forma normal. Se definieron cuatro tablas maestras (CLIENTES, "
        "SUCURSALES, FORMAS_PAGO y ARTICULOS), una tabla de encabezado (FACTURAS) y una tabla de detalle "
        "(DETALLE_FACTURA), que registra los artículos vendidos en cada factura y resuelve la relación de muchos a "
        "muchos entre facturas y artículos. Un cliente, una sucursal y una forma de pago pueden aparecer en muchas "
        "facturas; cada factura tiene una o varias líneas de detalle, y cada artículo puede venderse en muchas líneas.")
parrafo("Los valores monetarios usan el tipo NUMBER con dos decimales. El subtotal del detalle es una columna virtual "
        "que Oracle calcula como cantidad por valor unitario, y el detalle guarda su propio valor unitario para conservar "
        "el precio de la venta aunque cambie el precio del artículo.")
tabla("Tablas, llaves y tipos de datos principales",
      ["Tabla", "Llave primaria", "Llaves foráneas", "Tipos de datos principales"],
      [["CLIENTES", "id_cliente", "—", "NUMBER(6), VARCHAR2, DATE"],
       ["SUCURSALES", "id_sucursal", "—", "NUMBER(4), VARCHAR2"],
       ["FORMAS_PAGO", "id_forma_pago", "—", "NUMBER(2), VARCHAR2(40), NUMBER(3)"],
       ["ARTICULOS", "id_articulo", "—", "NUMBER(6), VARCHAR2, NUMBER(12,2), CHAR(1)"],
       ["FACTURAS", "numero_factura", "id_cliente, id_sucursal, id_forma_pago", "NUMBER(10), DATE, NUMBER(14,2)"],
       ["DETALLE_FACTURA", "numero_factura + numero_linea", "numero_factura, id_articulo", "NUMBER, NUMBER(12,2), NUMBER(14,2) virtual"]],
      [3.4, 3.8, 4.4, 4.4])

h2("Creación del usuario y de las tablas")
parrafo("En Oracle cada usuario es dueño de un esquema donde se guardan sus tablas; por eso, crear la base de datos del "
        "caso consistió en crear el usuario TECH_TRANSITION y, con su conexión, crear las seis tablas desde la Hoja de "
        "Trabajo de SQL Developer.")
figura("figura_17_crear_usuario.png", "Creación del usuario TECH_TRANSITION",
       "Desde la conexión del administrador (ADMIN_SYSTEM) se ejecuta CREATE USER para crear el usuario "
       "TECH_TRANSITION, propietario del esquema de facturación, con el tablespace USERS y una cuota de 200 MB. Luego, "
       "con GRANT se le otorgan los privilegios CREATE SESSION, CREATE TABLE, CREATE VIEW y CREATE SEQUENCE. La Salida "
       "de Script confirma «User TECH_TRANSITION creado.» y «Grant correcto.», por lo que el esquema del caso queda "
       "listo para recibir las tablas.")
figura("figura_18_conexion_tech.png", "Conexión con el usuario TECH_TRANSITION",
       "Se crea la conexión TECH_TRANSITION con ese mismo usuario, el host localhost, el puerto 1521 y el nombre de "
       "servicio FREEPDB1. Al presionar Probar, la parte inferior de la ventana muestra «Estado: Correcto», lo que "
       "confirma que el usuario del caso se comunica con la base de datos. Con esta conexión se desarrolla el resto "
       "del caso.")
figura("figura_19_tabla_clientes.png", "Creación de la tabla CLIENTES",
       "Con CREATE TABLE se crea la tabla CLIENTES, que guarda el tipo y número de documento, el nombre o razón "
       "social, el teléfono, el correo, la dirección, la ciudad y la fecha de registro de cada cliente. La columna "
       "id_cliente es la llave primaria (PK_CLIENTES), el documento no se puede repetir (UNIQUE) y el tipo de "
       "documento solo admite CC, CE o NIT (CHECK). La salida confirma «Table CLIENTES creado.».")
figura("figura_20_tabla_sucursales.png", "Creación de la tabla SUCURSALES",
       "Se crea la tabla SUCURSALES con las columnas id_sucursal, nombre_sucursal, ciudad, dirección y teléfono. La "
       "llave primaria es id_sucursal y el nombre de la sucursal es único. Esta tabla permite identificar qué sucursal "
       "emite cada factura. La salida muestra «Table SUCURSALES creado.».")
figura("figura_21_tabla_formas_pago.png", "Creación de la tabla FORMAS_PAGO",
       "Se crea la tabla FORMAS_PAGO con id_forma_pago como llave primaria, la descripción (única) y los días de "
       "plazo, que deben estar entre 0 y 180 (CHECK). Los días de plazo se usan para calcular la fecha de vencimiento "
       "de las facturas; 0 días corresponde a pago de contado. La salida confirma «Table FORMAS_PAGO creado.».")
figura("figura_22_tabla_articulos.png", "Creación de la tabla ARTICULOS",
       "Se crea la tabla ARTICULOS con el identificador (llave primaria), el código comercial (único), el nombre, la "
       "categoría, el precio unitario de tipo NUMBER(12,2) y el estado. Las restricciones CHECK exigen que el precio "
       "sea mayor que cero y que el estado sea A (activo) o I (inactivo). La salida muestra «Table ARTICULOS creado.».")
figura("figura_23_tabla_facturas.png", "Creación de la tabla FACTURAS",
       "Se crea la tabla FACTURAS, que almacena el número de factura (llave primaria), la fecha de factura, la fecha "
       "de vencimiento, el cliente, la sucursal, la forma de pago, el valor total y el valor pagado. Tres llaves "
       "foráneas la relacionan con CLIENTES, SUCURSALES y FORMAS_PAGO, y las restricciones CHECK impiden que el "
       "vencimiento sea anterior a la fecha de factura o que el valor pagado supere el valor total. La salida "
       "confirma «Table FACTURAS creado.».")
figura("figura_24_tabla_detalle_factura.png", "Creación de la tabla DETALLE_FACTURA",
       "Se crea la tabla DETALLE_FACTURA, que registra el artículo, la cantidad y el valor unitario de cada línea de "
       "la factura. Su llave primaria es compuesta (numero_factura y numero_linea), tiene llaves foráneas hacia "
       "FACTURAS y ARTICULOS, y la columna subtotal es virtual: Oracle la calcula como cantidad por valor unitario. "
       "La salida muestra «Table DETALLE_FACTURA creado.».")
figura("figura_25_indices_comentarios.png", "Creación de índices y comentarios",
       "Se ejecuta un script que crea cinco índices sobre las llaves foráneas y la fecha de factura, para agilizar "
       "las consultas, y agrega comentarios (COMMENT ON) que describen cada tabla y columna. En la Salida de Script "
       "se observan los mensajes «Comment creado.». En total el esquema quedó con 43 comentarios, que sirven como "
       "documentación dentro de la propia base de datos.")
figura("figura_26_tablas_y_columnas_facturas.png", "Tablas creadas y columnas de la tabla FACTURAS",
       "En el panel Conexiones, dentro de TECH_TRANSITION, la carpeta Tablas muestra las seis tablas creadas: "
       "ARTICULOS, CLIENTES, DETALLE_FACTURA, FACTURAS, FORMAS_PAGO y SUCURSALES. A la derecha, la pestaña Columnas "
       "de FACTURAS muestra el nombre, el tipo de dato, si admite nulos, el valor por defecto y el comentario de cada "
       "columna, lo que confirma que la estructura quedó registrada en Oracle.")

h2("Inserción de registros")
parrafo("Para comprobar el funcionamiento de la base de datos se insertaron datos de prueba ficticios con INSERT INTO "
        "y se confirmaron con COMMIT. Las facturas se registraron con valor total y valor pagado en cero; luego, el valor "
        "total se calculó a partir del detalle y se registraron los pagos con UPDATE.")
tabla("Registros insertados por tabla",
      ["Tabla", "Registros"],
      [["CLIENTES", "10"], ["SUCURSALES", "3"], ["FORMAS_PAGO", "6"], ["ARTICULOS", "15"],
       ["FACTURAS", "20"], ["DETALLE_FACTURA", "52"], ["Total", "106"]],
      [8.0, 4.0])
figura("figura_27_insertar_clientes.png", "Inserción de registros en CLIENTES",
       "Se ejecutan diez sentencias INSERT INTO para registrar los clientes, entre personas naturales (CC y CE) y "
       "empresas (NIT), y se confirma con COMMIT. La salida muestra «1 fila insertadas.» por cada registro y, al "
       "final, «Confirmación terminada.».")
figura("figura_28_consulta_clientes.png", "Consulta de los registros de CLIENTES",
       "Con SELECT * FROM clientes se consulta la tabla. La pestaña Resultado de la Consulta indica «Todas las Filas "
       "Recuperadas: 10» y muestra cada cliente con su tipo y número de documento, nombre, teléfono, correo, "
       "dirección, ciudad y fecha de registro.")
figura("figura_29_insertar_sucursales.png", "Inserción de registros en SUCURSALES",
       "Se registran las tres sucursales de la empresa: TT Bogotá Norte, TT Medellín El Poblado y TT Cali Granada, y "
       "se confirman con COMMIT. La salida muestra tres veces «1 fila insertadas.» y luego «Confirmación terminada.».")
figura("figura_30_consulta_sucursales.png", "Consulta de los registros de SUCURSALES",
       "La consulta devuelve las tres sucursales registradas con su ciudad, dirección y teléfono.")
figura("figura_31_insertar_formas_pago.png", "Inserción de registros en FORMAS_PAGO",
       "Se registran seis formas de pago: Efectivo, Tarjeta débito, Tarjeta crédito y Transferencia bancaria, con 0 "
       "días de plazo, y Crédito 30 días y Crédito 60 días. La salida confirma cada inserción y el COMMIT.")
figura("figura_32_consulta_formas_pago.png", "Consulta de los registros de FORMAS_PAGO",
       "La consulta muestra las seis formas de pago con sus días de plazo, que luego se usan para calcular la fecha "
       "de vencimiento de cada factura.")
figura("figura_33_insertar_articulos.png", "Inserción de registros en ARTICULOS",
       "Se registran quince artículos de tecnología (portátiles, monitores, periféricos, impresoras, consumibles, "
       "equipos de red, almacenamiento, energía y el servicio de instalación) con su código, categoría y precio. La "
       "salida muestra «1 fila insertadas.» por cada artículo y «Confirmación terminada.».")
figura("figura_34_consulta_articulos.png", "Consulta de los registros de ARTICULOS",
       "La consulta indica «Todas las Filas Recuperadas: 15» y muestra el código, el nombre, la categoría, el precio "
       "unitario y el estado activo (A) de cada artículo.")
figura("figura_35_insertar_facturas.png", "Inserción de registros en FACTURAS",
       "Se registran veinte facturas, numeradas de la 1001 a la 1020, con fechas entre julio y septiembre de 2026. "
       "Cada una indica el cliente, la sucursal y la forma de pago, y su fecha de vencimiento corresponde a la fecha "
       "de la factura más los días de plazo de la forma de pago. El valor total y el valor pagado inician en cero.")
figura("figura_36_insertar_detalle.png", "Inserción de registros en DETALLE_FACTURA",
       "Se registran 52 líneas de detalle, de dos a cuatro artículos por factura, con el número de factura, el número "
       "de línea, el artículo, la cantidad y el valor unitario. La salida confirma cada inserción y el COMMIT final.")
figura("figura_37_calcular_totales.png", "Cálculo del valor total y registro de pagos",
       "Con un UPDATE correlacionado se calcula el valor total de cada factura como la suma de los subtotales de su "
       "detalle, y la salida indica «20 filas actualizadas.». Después se registra el pago total de 15 facturas («15 "
       "filas actualizadas.») y el pago del 50 % de 3 facturas a crédito («3 filas actualizadas.»); las 2 restantes "
       "quedan sin pago. Los cambios se confirman con COMMIT.")
figura("figura_38_consulta_facturas.png", "Consulta de los registros de FACTURAS",
       "La consulta devuelve las 20 facturas con su número, fechas, cliente, sucursal, forma de pago, valor total y "
       "valor pagado. Por ejemplo, la factura 1001 tiene un valor total de $19.725.000 pagado en su totalidad, "
       "mientras que las facturas 1016 y 1020 aparecen con valor pagado en cero.")
figura("figura_39_consulta_detalle.png", "Consulta de los registros de DETALLE_FACTURA",
       "La consulta muestra las líneas de detalle con la cantidad, el valor unitario y el subtotal calculado "
       "automáticamente por Oracle. SQL Developer indica «Se han recuperado 50 filas» porque trae los resultados en "
       "bloques de 50; la tabla contiene 52 registros, como se comprobó con el conteo de la Tabla 4.")

h2("Relaciones entre tablas")
parrafo("Las relaciones se implementaron con llaves primarias, que identifican cada registro de forma única, y llaves "
        "foráneas, que obligan a que el valor referenciado exista en la tabla relacionada.")
tabla("Relaciones implementadas",
      ["Llave foránea", "Tabla hija (columna)", "Tabla padre (llave primaria)", "Cardinalidad"],
      [["FK_FACTURAS_CLIENTES", "FACTURAS (id_cliente)", "CLIENTES (id_cliente)", "1 : N"],
       ["FK_FACTURAS_SUCURSALES", "FACTURAS (id_sucursal)", "SUCURSALES (id_sucursal)", "1 : N"],
       ["FK_FACTURAS_FORMAS_PAGO", "FACTURAS (id_forma_pago)", "FORMAS_PAGO (id_forma_pago)", "1 : N"],
       ["FK_DETALLE_FACTURAS", "DETALLE_FACTURA (numero_factura)", "FACTURAS (numero_factura)", "1 : N"],
       ["FK_DETALLE_ARTICULOS", "DETALLE_FACTURA (id_articulo)", "ARTICULOS (id_articulo)", "1 : N"]],
      [4.6, 4.6, 4.6, 2.2])
figura("figura_40_modelo_relaciones.png", "Modelo de relaciones de la tabla FACTURAS",
       "En la pestaña Model de la tabla FACTURAS, SQL Developer lee las llaves foráneas definidas en Oracle y dibuja "
       "las tablas relacionadas: FACTURAS se vincula con CLIENTES, SUCURSALES y FORMAS_PAGO, y DETALLE_FACTURA se "
       "vincula con FACTURAS. Esta vista demuestra que las relaciones existen en la base de datos y no solo en el "
       "diseño.")
figura("figura_41_claves.png", "Consulta de las llaves primarias y foráneas",
       "Se consultan las vistas USER_CONSTRAINTS y USER_CONS_COLUMNS para listar las llaves del esquema. El "
       "resultado (12 filas) muestra la llave primaria de cada tabla, incluida la llave compuesta de DETALLE_FACTURA, "
       "y las cinco llaves foráneas con su tabla referenciada, entre ellas FK_DETALLE_ARTICULOS hacia ARTICULOS.")

h2("Consultas de comprobación")
parrafo("Para comprobar que la base de datos responde a las necesidades del departamento de facturación se ejecutaron "
        "consultas con JOIN, GROUP BY, SUM, COUNT y ORDER BY. La Tabla 6 resume los resultados.")
tabla("Resumen de las consultas de comprobación",
      ["Consulta", "Resultado obtenido"],
      [["Facturas con cliente, sucursal y forma de pago (JOIN)", "20 filas"],
       ["Detalle de facturas con artículo (JOIN)", "8 líneas de las facturas 1001 a 1003"],
       ["Facturación por sucursal", "Bogotá Norte: 8 facturas, $86.950.000; Medellín El Poblado: 6, $67.720.000; Cali Granada: 6, $17.745.000"],
       ["Facturación por forma de pago", "Crédito 30 días es la más usada: 6 facturas, $93.245.000"],
       ["Totales", "Facturado $172.415.000; pagado $133.580.000; saldo $38.835.000; recaudo 77,48 %"],
       ["Saldos pendientes", "5 facturas con saldo; la factura 1009 está vencida"],
       ["Artículos más vendidos", "Memoria USB 64 GB (38 unidades) y Portátil 14 pulgadas Core i5 (32 unidades)"]],
      [6.5, 9.5])
figura("figura_42_join_facturas.png", "Consulta de facturas con cliente, sucursal y forma de pago",
       "Se ejecuta una consulta que une con JOIN las tablas FACTURAS, CLIENTES, SUCURSALES y FORMAS_PAGO. El "
       "resultado (20 filas) muestra, en lugar de códigos, el nombre del cliente, la sucursal y la forma de pago de "
       "cada factura; por ejemplo, la factura 1001 de Comercializadora Andina Delta S.A.S., emitida en TT Bogotá Norte "
       "con Crédito 30 días. Esto comprueba que las relaciones permiten reconstruir la información completa.")
figura("figura_43_detalle_join.png", "Consulta del detalle de facturas con el nombre del artículo",
       "Se une DETALLE_FACTURA con ARTICULOS para ver qué se vendió en las facturas 1001, 1002 y 1003. El resultado "
       "muestra 8 líneas con el código y nombre del artículo, la cantidad, el valor unitario y el subtotal; por "
       "ejemplo, 5 portátiles de $3.200.000 en la factura 1001, con un subtotal de $16.000.000.")
figura("figura_44_por_sucursal.png", "Facturación por sucursal",
       "Con GROUP BY, COUNT y SUM se agrupan las facturas por sucursal. TT Bogotá Norte registra 8 facturas por "
       "$86.950.000, TT Medellín El Poblado 6 por $67.720.000 y TT Cali Granada 6 por $17.745.000.")
figura("figura_45_por_forma_pago.png", "Facturación por forma de pago",
       "Se agrupan las facturas por forma de pago con LEFT JOIN, para incluir todas las formas registradas. Crédito "
       "30 días es la más utilizada (6 facturas por $93.245.000), seguida de Crédito 60 días (3 facturas por "
       "$62.610.000); Efectivo es la menos usada (2 facturas por $725.000).")
figura("figura_46_totales.png", "Total facturado, total pagado y saldo pendiente",
       "La consulta calcula los totales de las 20 facturas: $172.415.000 facturados, $133.580.000 pagados y "
       "$38.835.000 de saldo pendiente, lo que equivale a un recaudo del 77,48 %.")
figura("figura_47_saldos.png", "Facturas con saldo pendiente",
       "Se consultan las facturas cuyo valor pagado es menor que el valor total. Cinco facturas tienen saldo "
       "pendiente; con CASE se marca como VENCIDA la factura 1009 de Distribuidora Café del Valle S.A.S., que venció "
       "el 2 de septiembre de 2026 con un saldo de $5.685.000, y las demás aparecen AL DIA.")
figura("figura_48_mas_vendidos.png", "Artículos más vendidos",
       "Se agrupa el detalle por artículo y se ordena por unidades vendidas. El más vendido es la Memoria USB 64 GB "
       "(38 unidades), seguido del Portátil 14 pulgadas Core i5 (32 unidades), que además genera el mayor valor de "
       "ventas ($102.400.000).")

h2("Actualización y eliminación de registros")
parrafo("Por último, se practicaron las operaciones UPDATE y DELETE sobre los datos del caso y se comprobó que Oracle "
        "rechaza las operaciones que violan la integridad de la información.")
figura("figura_49_update.png", "Registro de un pago con UPDATE",
       "Con UPDATE facturas SET valor_pagado = valor_total WHERE numero_factura = 1009 se registra el pago del saldo "
       "pendiente de la factura vencida. La salida confirma «1 fila actualizadas.» y «Confirmación terminada.» tras "
       "el COMMIT.")
figura("figura_50_consulta_update.png", "Consulta de la factura 1009 después del pago",
       "La consulta muestra que la factura 1009 tiene un valor total de $11.370.000 y un valor pagado de $11.370.000, "
       "por lo que su saldo quedó en cero.")
figura("figura_51_delete.png", "Inserción y eliminación de un registro de prueba con DELETE",
       "Se agrega la forma de pago de prueba «Cheque (prueba)» con el identificador 7 y luego se elimina con DELETE "
       "FROM formas_pago WHERE id_forma_pago = 7. La salida muestra «1 fila insertadas.», «1 fila eliminado» y "
       "«Confirmación terminada.». Con esto se practicaron las operaciones de insertar, consultar, actualizar y "
       "eliminar.")
figura("figura_52_integridad.png", "Pruebas de integridad de los datos",
       "Se intentan tres operaciones inválidas. Oracle rechaza la eliminación de un cliente que tiene facturas "
       "(ORA-02292, registro secundario encontrado), la factura de un cliente que no existe (ORA-02291, clave "
       "principal no encontrada) y la línea de detalle con cantidad negativa (ORA-02290, restricción de control "
       "CK_DETALLE_CANTIDAD). Estas pruebas demuestran que las llaves y restricciones protegen la información de "
       "facturación.")

h1("Conclusiones del caso de estudio")
parrafo("El caso TECH TRANSITION permitió convertir una información de facturación dispersa y sin control en una base "
        "de datos relacional en Oracle, organizada en seis tablas: CLIENTES, SUCURSALES, FORMAS_PAGO, ARTICULOS, "
        "FACTURAS y DETALLE_FACTURA. Todos los datos exigidos por la guía (número de factura, fechas de factura y de "
        "vencimiento, cliente, sucursal, valor total, forma de pago, valor pagado, artículo, cantidad y valor unitario) "
        "quedaron almacenados una sola vez y relacionados mediante llaves primarias y foráneas.")
parrafo("La lista de chequeo confirmó que se cumplieron los requerimientos del caso. Las restricciones de integridad "
        "rechazaron los datos inválidos, algo que una hoja de cálculo no garantiza, y las consultas con JOIN, GROUP BY "
        "y SUM transformaron los 106 registros de prueba en información útil para la toma de decisiones: un total "
        "facturado de $172.415.000 con un recaudo del 77,48 %, la identificación de una factura vencida y de los "
        "artículos más vendidos. Esta base de datos es el punto de partida para la futura bodega de datos del "
        "departamento de facturación.")

doc.save(SALIDA)
print("Generado:", SALIDA, "| figuras", 17, "a", fig_num[0] - 1, "| tablas", tab_num[0] - 1)
